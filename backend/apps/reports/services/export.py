import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse

class ExportService:
    @staticmethod
    def export_csv(records):
        """
        Generates a CSV response for attendance records.
        """
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            'Employee Code', 'Employee Name', 'Department', 'Date',
            'Check In', 'Check Out', 'Working Minutes', 'Working Hours',
            'Status', 'IP Verified', 'Manual Correction'
        ])

        for r in records:
            working_hours = round(r.working_minutes / 60.0, 2) if r.working_minutes else 0.0
            writer.writerow([
                r.employee.employee_code,
                r.employee.full_name,
                r.employee.department.name if r.employee.department else 'N/A',
                str(r.attendance_date),
                r.check_in_at.strftime('%Y-%m-%d %H:%M:%S') if r.check_in_at else '',
                r.check_out_at.strftime('%Y-%m-%d %H:%M:%S') if r.check_out_at else '',
                r.working_minutes,
                working_hours,
                r.attendance_status,
                r.verification_result,
                'Yes' if r.is_manual_correction else 'No'
            ])

        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="attendance_report.csv"'
        return response

    @staticmethod
    def export_excel(records):
        """
        Generates an Excel (.xlsx) file using openpyxl with styled headers.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Report"

        headers = [
            'Employee Code', 'Employee Name', 'Department', 'Date',
            'Check In', 'Check Out', 'Working Minutes', 'Working Hours',
            'Status', 'IP Verified', 'Manual Correction'
        ]
        ws.append(headers)

        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for r in records:
            working_hours = round(r.working_minutes / 60.0, 2) if r.working_minutes else 0.0
            ws.append([
                r.employee.employee_code,
                r.employee.full_name,
                r.employee.department.name if r.employee.department else 'N/A',
                str(r.attendance_date),
                r.check_in_at.strftime('%Y-%m-%d %H:%M:%S') if r.check_in_at else '',
                r.check_out_at.strftime('%Y-%m-%d %H:%M:%S') if r.check_out_at else '',
                r.working_minutes,
                working_hours,
                r.attendance_status,
                r.verification_result,
                'Yes' if r.is_manual_correction else 'No'
            ])

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="attendance_report.xlsx"'
        return response
